import json
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)
from docx import Document as WordDocument
from openpyxl import Workbook

from chat.models import Conversation, Message
from company.models import Company, CompanyMembership
from docs.models import Document, DocumentChunk
from docs.services import (
    DocumentValidationError,
    extract_document_chunks,
    split_page_text,
)
from docs.tasks import process_document_task
from infrastructure.vectorstore.chroma_backend import ChromaVectorStore


class PdfChunkingTests(SimpleTestCase):
    def test_long_pages_are_split_with_overlap_and_page_references(self):
        text = "A" * 2500

        chunks = split_page_text(text, page_number=4)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(page_number == 4 for page_number, _ in chunks))
        self.assertEqual(chunks[0][1][-180:], chunks[1][1][:180])


def make_pdf(text):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font_ref}
            )
        }
    )
    content = DecodedStreamObject()
    safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content.set_data(
        f"BT /F1 12 Tf 50 720 Td ({safe_text}) Tj ET".encode("latin-1")
    )
    page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


class TestOpenAIProvider:
    model = "test-model"

    def __init__(self):
        self.prompts = []
        self.embedded_texts = []
        self.embedding_batch_sizes = []

    def embed_documents(self, texts):
        self.embedded_texts.extend(texts)
        self.embedding_batch_sizes.append(len(texts))
        return [[1.0, 0.0, 1.0] for _ in texts]

    def embed_query(self, text):
        return [1.0, 0.0, 1.0]

    def generate_structured(self, prompt, schema):
        self.prompts.append(prompt)
        return schema(
            response="The document says records are retained for seven years.",
            chart=None,
        )


def make_docx():
    document = WordDocument()
    document.add_paragraph("Quarterly revenue by region")
    document.add_table(rows=2, cols=2)
    document.tables[0].cell(0, 0).text = "Region"
    document.tables[0].cell(0, 1).text = "Revenue"
    document.tables[0].cell(1, 0).text = "West"
    document.tables[0].cell(1, 1).text = "125"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Month", "Revenue"])
    sheet.append(["January", 120])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class PdfToChatFlowTests(TestCase):
    def setUp(self):
        self.temp_dir = TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        base_path = Path(self.temp_dir.name)
        self.settings_override = override_settings(
            MEDIA_ROOT=base_path / "media",
            CHROMA_HOST="",
            CHROMA_PATH=base_path / "chroma",
        )
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)

        self.user = get_user_model().objects.create_user(
            username="pdfuser",
            email="pdfuser@example.com",
            password="secret123",
        )
        self.client = Client()
        self.client.force_login(self.user)
        self.provider = TestOpenAIProvider()
        self.document_provider_patch = patch(
            "docs.services.get_ai_provider",
            return_value=self.provider,
        )
        self.chat_provider_patch = patch(
            "chat.services.get_ai_provider",
            return_value=self.provider,
        )
        self.document_provider_patch.start()
        self.chat_provider_patch.start()
        self.addCleanup(self.document_provider_patch.stop)
        self.addCleanup(self.chat_provider_patch.stop)
        self.dispatch_patch = patch(
            "docs.views.dispatch_document_processing",
            return_value=("test-task-id", "celery"),
        )
        self.dispatch_patch.start()
        self.addCleanup(self.dispatch_patch.stop)

    def upload_pdf(self, name="records-policy.pdf", text=None):
        pdf = make_pdf(
            text
            or "Records are retained for seven years before secure disposal."
        )
        return self.client.post(
            reverse("document_upload"),
            {
                "file": SimpleUploadedFile(
                    name,
                    pdf,
                    content_type="application/pdf",
                )
            },
        )

    def test_pdf_upload_chunks_indexes_and_answers_with_database_references(self):
        upload_response = self.upload_pdf()
        self.assertEqual(upload_response.status_code, 202, upload_response.content)
        payload = upload_response.json()["document"]
        document = Document.objects.get(pk=payload["id"])
        self.assertEqual(document.status, Document.STATUS_QUEUED)
        process_document_task(str(document.pk))
        document.refresh_from_db()
        chunk = DocumentChunk.objects.get(document=document)
        conversation = Conversation.objects.get(
            pk=upload_response.json()["conversation_id"],
            user=self.user,
        )
        attachment_message = conversation.messages.get()

        self.assertEqual(document.status, Document.STATUS_READY)
        self.assertEqual(document.processing_stage, "ready")
        self.assertEqual(document.processing_progress, 100)
        self.assertEqual(document.chunk_count, 1)
        self.assertIn("seven years", chunk.content)
        self.assertEqual(payload["reference_id"], str(document.pk))
        self.assertEqual(document.conversation, conversation)
        self.assertEqual(list(attachment_message.documents.all()), [document])
        self.assertEqual(
            attachment_message.content,
            "Attached file: records-policy.pdf",
        )

        history_response = self.client.get(
            reverse(
                "conversation_detail",
                kwargs={"conversation_id": conversation.pk},
            )
        )
        self.assertEqual(history_response.status_code, 200)
        self.assertEqual(
            history_response.json()["messages"][0]["documents"][0]["reference_id"],
            str(document.pk),
        )
        self.assertTrue(
            history_response.json()["messages"][0]["documents"][0]["file_url"]
        )

        stored_vectors = ChromaVectorStore().similarity_search(
            [1.0, 0.0, 1.0],
            collection="pdf_chunks",
            filter={"user_id": str(self.user.pk)},
        )
        self.assertEqual(stored_vectors[0]["id"], str(chunk.pk))
        self.assertEqual(stored_vectors[0]["metadata"]["document_id"], str(document.pk))
        stored_embeddings = ChromaVectorStore().get_embeddings(
            [str(chunk.pk)],
            collection="pdf_chunks",
        )
        self.assertEqual(stored_embeddings[str(chunk.pk)], [1.0, 0.0, 1.0])

        chat_response = self.client.post(
            reverse("chat_api"),
            {
                "message": "How long are records retained?",
                "conversation_id": str(conversation.pk),
            },
            content_type="application/json",
        )
        self.assertEqual(chat_response.status_code, 200, chat_response.content)
        result = chat_response.json()
        self.assertEqual(result["conversation_id"], str(conversation.pk))
        self.assertEqual(Conversation.objects.count(), 1)
        self.assertEqual(conversation.messages.count(), 3)
        self.assertIn(str(chunk.pk), self.provider.prompts[-1])
        self.assertEqual(result["citations"][0]["reference_id"], str(chunk.pk))
        self.assertEqual(result["citations"][0]["document_id"], str(document.pk))
        self.assertEqual(result["citations"][0]["page"], 1)
        self.assertEqual(Message.objects.get(role=Message.ROLE_ASSISTANT).citations, result["citations"])
        file_response = self.client.get(
            reverse(
                "document_file",
                kwargs={"file_path": document.file.name},
            )
        )
        self.assertEqual(file_response.status_code, 200)

    def test_uploaded_files_are_assigned_to_the_active_company(self):
        company = Company.objects.create(name="Acme", slug="acme")
        CompanyMembership.objects.create(
            user=self.user,
            company=company,
            role=CompanyMembership.Role.MEMBER,
        )

        response = self.upload_pdf("company-policy.pdf")

        self.assertEqual(response.status_code, 202, response.content)
        document = Document.objects.get(pk=response.json()["document"]["id"])
        conversation = Conversation.objects.get(pk=response.json()["conversation_id"])
        self.assertEqual(document.company, company)
        self.assertEqual(conversation.company, company)

    def test_upload_saves_the_question_as_the_attachment_message(self):
        question = "Explain this CSV and create charts"
        response = self.client.post(
            reverse("document_upload"),
            {
                "file": SimpleUploadedFile(
                    "education.csv",
                    b"country,count\nA,10\nB,20\n",
                    content_type="text/csv",
                ),
                "message": question,
            },
        )

        self.assertEqual(response.status_code, 202, response.content)
        message = Message.objects.get(
            conversation_id=response.json()["conversation_id"],
        )
        self.assertEqual(message.content, question)
        self.assertEqual(response.json()["message"]["id"], str(message.pk))
        self.assertEqual(
            [document.title for document in message.documents.all()],
            ["education.csv"],
        )

    def test_rejects_non_pdf_and_invalid_pdf_content(self):
        not_pdf = self.client.post(
            reverse("document_upload"),
            {
                "file": SimpleUploadedFile(
                    "not-a-pdf.pdf",
                    b"not a PDF",
                    content_type="application/pdf",
                )
            },
        )
        self.assertEqual(not_pdf.status_code, 400)

        unreadable_pdf = self.client.post(
            reverse("document_upload"),
            {
                "file": SimpleUploadedFile(
                    "broken.pdf",
                    b"%PDF-1.4 broken content",
                    content_type="application/pdf",
                )
            },
        )
        self.assertEqual(unreadable_pdf.status_code, 202)
        document = Document.objects.get(title="broken.pdf")
        with self.assertRaises(DocumentValidationError):
            process_document_task(str(document.pk))
        document.refresh_from_db()
        self.assertEqual(document.status, Document.STATUS_FAILED)

    def test_uploads_supported_text_tabular_and_office_formats(self):
        uploads = [
            (
                "notes.txt",
                b"Keep records for seven years.",
                "text/plain",
                "seven years",
            ),
            (
                "summary.md",
                b"# Results\nRevenue increased by 15 percent.",
                "text/markdown",
                "15 percent",
            ),
            (
                "sales.csv",
                b"Month,Revenue\nJanuary,120\nFebruary,150\n",
                "text/csv",
                "January | 120",
            ),
            (
                "report.json",
                json.dumps({"region": "West", "revenue": 125}).encode(),
                "application/json",
                '"revenue": 125',
            ),
            (
                "report.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "West | 125",
            ),
            (
                "report.xlsx",
                make_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "January | 120",
            ),
        ]

        for name, content, content_type, expected_text in uploads:
            with self.subTest(name=name):
                response = self.client.post(
                    reverse("document_upload"),
                    {
                        "file": SimpleUploadedFile(name, content, content_type=content_type)
                    },
                )
                self.assertEqual(response.status_code, 202, response.content)
                document = Document.objects.get(pk=response.json()["document"]["id"])
                process_document_task(str(document.pk))
                document.refresh_from_db()
                chunk = DocumentChunk.objects.get(document=document)
                self.assertEqual(document.status, Document.STATUS_READY)
                self.assertEqual(document.file_type, name.rsplit(".", 1)[-1])
                self.assertEqual(document.content_type, content_type)
                self.assertEqual(document.file_size, len(content))
                self.assertIn(expected_text, chunk.content)
                self.assertTrue(chunk.source_location)

    def test_large_csv_chunks_compactly_with_headers_and_row_locations(self):
        headers = [f"measure_{index:02}_long_header" for index in range(40)]
        output = [",".join(headers)]
        for row_number in range(10_001):
            output.append(
                ",".join(f"value_{row_number:05}_{column:02}" for column in range(40))
            )
        csv_content = ("\n".join(output) + "\n").encode()
        document = Document(
            title="large.csv",
            file=SimpleUploadedFile(
                "large.csv",
                csv_content,
                content_type="text/csv",
            ),
            file_type="csv",
        )

        chunks = extract_document_chunks(document)

        self.assertLess(len(chunks), 1_000)
        self.assertGreater(len(chunks), 100)
        self.assertIn("measure_00_long_header", chunks[0][2])
        self.assertIn("value_00000_00", chunks[0][2])
        self.assertTrue(chunks[-1][1].endswith("10002"))
        self.assertTrue(any("value_10000_39" in content for _, _, content in chunks))

    def test_csv_indexer_skips_leading_metadata_comment(self):
        document = Document(
            title="people.csv",
            file=SimpleUploadedFile(
                "people.csv",
                b"# Sample file metadata\nID,Age\n1,18\n2,22\n",
                content_type="text/csv",
            ),
            file_type="csv",
        )

        chunks = extract_document_chunks(document)

        self.assertEqual(len(chunks), 1)
        self.assertIn("Columns: ID | Age", chunks[0][2])
        self.assertIn("3: 1 | 18", chunks[0][2])
        self.assertNotIn("Sample file metadata", chunks[0][2])

    def test_rejects_unsupported_extension_and_invalid_json(self):
        missing_file = self.client.post(reverse("document_upload"))
        self.assertEqual(missing_file.status_code, 400)
        self.assertIn("supported file", missing_file.json()["error"].lower())

        unsupported = self.client.post(
            reverse("document_upload"),
            {
                "file": SimpleUploadedFile(
                    "archive.zip",
                    b"PK\x03\x04archive",
                    content_type="application/zip",
                )
            },
        )
        self.assertEqual(unsupported.status_code, 400)
        self.assertIn("Supported files", unsupported.json()["error"])

        invalid_json = self.client.post(
            reverse("document_upload"),
            {
                "file": SimpleUploadedFile(
                    "broken.json",
                    b'{"broken":',
                    content_type="application/json",
                )
            },
        )
        self.assertEqual(invalid_json.status_code, 202)
        document = Document.objects.get(title="broken.json")
        with self.assertRaises(DocumentValidationError):
            process_document_task(str(document.pk))
        document.refresh_from_db()
        self.assertEqual(document.status, Document.STATUS_FAILED)

    @patch("docs.services.ChromaVectorStore.add_documents")
    def test_failed_document_can_be_reprocessed_without_uploading_again(self, add_documents):
        add_documents.side_effect = lambda documents, collection: [
            str(item["id"]) for item in documents
        ]
        document = Document.objects.create(
            uploaded_by=self.user,
            title="retry-notes.txt",
            file=SimpleUploadedFile(
                "retry-notes.txt",
                b"Reprocessing should keep the original media file.",
                content_type="text/plain",
            ),
            file_type="txt",
            content_type="text/plain",
            file_size=48,
            status=Document.STATUS_FAILED,
            processing_error="The previous chunk limit was reached.",
        )

        response = self.client.post(
            reverse("document_reprocess", kwargs={"document_id": document.pk})
        )

        self.assertEqual(response.status_code, 202, response.content)
        document.refresh_from_db()
        self.assertEqual(document.status, Document.STATUS_QUEUED)
        process_document_task(str(document.pk))
        document.refresh_from_db()
        self.assertEqual(document.status, Document.STATUS_READY)
        self.assertEqual(document.chunk_count, 1)
        self.assertEqual(response.json()["document"]["id"], str(document.pk))

    def test_legacy_queued_document_without_task_id_can_be_retried(self):
        document = Document.objects.create(
            uploaded_by=self.user,
            title="stuck-queued.txt",
            file=SimpleUploadedFile(
                "stuck-queued.txt",
                b"Retry a queued document from before background tasks were enabled.",
                content_type="text/plain",
            ),
            file_type="txt",
            status=Document.STATUS_QUEUED,
        )

        response = self.client.post(
            reverse("document_reprocess", kwargs={"document_id": document.pk})
        )

        self.assertEqual(response.status_code, 202, response.content)
        document.refresh_from_db()
        self.assertEqual(document.processing_task_id, "test-task-id")

    @patch("docs.services.extract_document_chunks")
    def test_upload_queues_without_extracting_in_the_request(self, extract_chunks):
        response = self.upload_pdf()

        self.assertEqual(response.status_code, 202, response.content)
        extract_chunks.assert_not_called()
        document = Document.objects.get(pk=response.json()["document"]["id"])
        self.assertEqual(document.status, Document.STATUS_QUEUED)
        self.assertEqual(document.processing_stage, "queued")
        self.assertEqual(document.processing_task_id, "test-task-id")

    @override_settings(DEBUG=True)
    @patch("docs.tasks.process_document_task.apply_async")
    @patch("docs.tasks._local_executor.submit")
    def test_local_upload_starts_background_processing_without_redis(
        self,
        submit,
        apply_async,
    ):
        self.dispatch_patch.stop()

        response = self.upload_pdf()

        self.assertEqual(response.status_code, 202, response.content)
        document = Document.objects.get(pk=response.json()["document"]["id"])
        self.assertEqual(document.status, Document.STATUS_QUEUED)
        self.assertEqual(document.processing_backend, "local-background")
        self.assertTrue(document.processing_task_id)
        submit.assert_called_once()
        apply_async.assert_not_called()

    @patch(
        "docs.services.ChromaVectorStore.add_documents",
        side_effect=lambda items, collection: [item["id"] for item in items],
    )
    @patch("docs.services.extract_document_chunks")
    def test_worker_batches_embeddings_and_persists_completion(
        self,
        extract_chunks,
        _add_documents,
    ):
        extract_chunks.return_value = [
            (1, f"row {index}", f"Revenue row {index}")
            for index in range(600)
        ]
        document = Document.objects.create(
            uploaded_by=self.user,
            title="many-rows.txt",
            file=SimpleUploadedFile(
                "many-rows.txt",
                b"indexed by a background task",
                content_type="text/plain",
            ),
            file_type="txt",
            content_type="text/plain",
            file_size=29,
            status=Document.STATUS_QUEUED,
        )

        process_document_task(str(document.pk))

        document.refresh_from_db()
        self.assertEqual(
            sorted(self.provider.embedding_batch_sizes),
            [24, 96, 96, 96, 96, 96, 96],
        )
        self.assertEqual(document.status, Document.STATUS_READY)
        self.assertEqual(document.chunk_count, 600)
        self.assertEqual(document.processing_progress, 100)

    @override_settings(DEBUG=False)
    @patch(
        "docs.views.dispatch_document_processing",
        side_effect=ConnectionError("broker unavailable"),
    )
    def test_broker_failure_is_persisted_and_returned(self, _dispatch):
        response = self.upload_pdf()

        self.assertEqual(response.status_code, 503)
        document = Document.objects.get(pk=response.json()["document"]["id"])
        self.assertEqual(document.status, Document.STATUS_FAILED)
        self.assertIn("Redis and a Celery worker", document.processing_error)

    @override_settings(DEBUG=True)
    @patch("docs.tasks.process_document_task.apply_async")
    @patch("docs.tasks._local_executor.submit")
    def test_local_debug_dispatch_falls_back_to_background_executor(
        self,
        submit,
        apply_async,
    ):
        from docs.tasks import dispatch_document_processing

        task_id, backend = dispatch_document_processing("document-id")

        self.assertTrue(task_id)
        self.assertEqual(backend, "local-background")
        apply_async.assert_not_called()
        submit.assert_called_once()
        self.assertEqual(submit.call_args.args[1], "document-id")
        self.assertEqual(
            submit.call_args.kwargs,
            {"backend": "local-background"},
        )

    def test_documents_require_login_and_are_private_to_owner(self):
        upload_response = self.upload_pdf()
        document = Document.objects.get(pk=upload_response.json()["document"]["id"])
        other_client = Client()

        self.assertEqual(
            other_client.get(reverse("document_list")).status_code,
            401,
        )
        self.assertEqual(
            other_client.get(
                reverse(
                    "document_file",
                    kwargs={"file_path": document.file.name},
                )
            ).status_code,
            401,
        )

        other_user = get_user_model().objects.create_user(username="other")
        other_client.force_login(other_user)
        self.assertEqual(
            other_client.get(
                reverse(
                    "document_file",
                    kwargs={"file_path": document.file.name},
                )
            ).status_code,
            404,
        )
